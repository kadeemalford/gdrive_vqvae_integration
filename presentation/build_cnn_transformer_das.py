"""Build an editable presentation reviewing the paper and its DAS adaptation.

Requires python-pptx. Run from any directory with the project virtualenv.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "CNN_TRANSFORMER_E_RECONSTRUCAO_DAS.pptx"
NAVY = "102238"
TEAL = "007F86"
BLUE = "2864B4"
GOLD = "B56B13"
INK = "24364A"
MUTED = "52677B"
LIGHT = "EEF4F8"
WHITE = "FFFFFF"
PAPER = "Estevão e Rocha · CBA_2026_paper__7079.pdf"
REPORT = "Análise do projeto · relatório 60 e PROGRESS_CHECKPOINT_1"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = "CNN–Transformer: do reconhecimento de equipamentos à reconstrução DAS"
prs.core_properties.subject = "Avaliação crítica do artigo e proposta de comparação neural FEBUS"
prs.core_properties.language = "pt-BR"


def rect(slide, x, y, w, h, color, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    return shape


def text(slide, x, y, w, h, content, size=21, color=INK, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    for i, line in enumerate(content.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(10)
    return box


def base(section, title, subtitle="", source=PAPER, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    rect(slide, 0, 0, 0.16, 7.5, TEAL)
    text(slide, .55, .26, 12, .35, section.upper(), 12, TEAL, True)
    text(slide, .55, .82, 12.2, .9, title, 26 if len(title) > 68 else 30, NAVY, True)
    if subtitle:
        text(slide, .55, 1.82, 12.1, .55, subtitle, 17, MUTED)
    rect(slide, .55, 7.02, 12.2, .015, "D8E3EC")
    text(slide, .55, 7.12, 11.6, .23, source, 10, MUTED)
    text(slide, 12.05, 7.1, .7, .3, f"{len(prs.slides):02d}", 12, MUTED)
    slide.notes_slide.notes_text_frame.text = notes or source
    return slide


def card(slide, x, y, w, h, heading, body, color=TEAL, size=21):
    rect(slide, x, y, w, h, LIGHT, True)
    rect(slide, x, y, .055, h, color)
    text(slide, x+.2, y+.17, w-.4, .55, heading, 22, color, True)
    box = text(slide, x+.2, y+.88, w-.4, h-1, body, size)
    for paragraph in box.text_frame.paragraphs:
        paragraph.space_after = Pt(6)


def banner(slide, message, y=6.15, color=NAVY):
    rect(slide, .55, y, 12.2, .64, color, True)
    text(slide, .78, y+.06, 11.75, .56, message, 16, WHITE, True)


def node(slide, x, y, w, h, title, color=TEAL, size=19):
    rect(slide, x, y, w, h, color, True)
    box = text(slide, x+.13, y+.08, w-.26, h-.1, title, min(size, 16), WHITE, True)
    for paragraph in box.text_frame.paragraphs:
        paragraph.space_after = Pt(0)


def link(slide, x1, y1, x2, y2):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = RGBColor.from_string(MUTED)
    shape.line.width = Pt(2)


def table(slide, headers, rows, widths, y=2.5, row_h=.65, font=19):
    x = .65
    for header, width in zip(headers, widths):
        rect(slide, x, y, width, row_h, NAVY)
        text(slide, x+.12, y+.12, width-.2, row_h-.15, header, font, WHITE, True)
        x += width
    for i, row in enumerate(rows):
        x = .65
        for value, width in zip(row, widths):
            rect(slide, x, y+(i+1)*row_h, width, row_h, LIGHT if i % 2 == 0 else "F9FBFD")
            text(slide, x+.12, y+(i+1)*row_h+.12, width-.2, row_h-.15, value, font)
            x += width


s = base("Artigo e proposta de adaptação", "CNN–Transformer e reconstrução DAS",
         "Classificação de equipamentos, aprendizado incremental e eficiência de representação neural",
         notes="Apresentação baseada na leitura integral do artigo de Arthur Coelho Estevão e Helder Roberto de Oliveira Rocha. As adaptações DAS são propostas do projeto, não resultados do artigo. Nenhum novo experimento A/B/C foi iniciado para esta apresentação.")
text(s, .65, 2.65, 11.9, 1.15, "Do reconhecimento de dispositivos elétricos\na um codec generalista para diferentes famílias de fibras", 29, NAVY, True)
card(s, .65, 4.35, 5.85, 1.5, "O que o artigo demonstra", "", BLUE)
card(s, 6.7, 4.35, 5.85, 1.5, "O que vale testar em DAS", "", TEAL)
text(s, .85, 5.1, 5.4, .5, "Identificação, reconstrução e expansão", 18)
text(s, 6.9, 5.1, 5.4, .5, "Representação canal–tempo com bits controlados", 18)
banner(s, "Contextualização  →  Arquitetura  →  Resultados  →  Aplicação em DAS")

s = base("01 · Contextualização", "O problema: reconhecer um catálogo que muda",
         "NILM = monitoramento não intrusivo de cargas elétricas.",
         source=PAPER+" · §§1–2, 4.1 e 5")
card(s, .65, 2.65, 3.85, 3.1, "Objetivo do domínio", "Inferir equipamentos e suas contribuições a partir de medições elétricas.", size=22)
card(s, 4.73, 2.65, 3.85, 3.1, "Limite closed-set", "Um dispositivo não visto pode ser atribuído incorretamente a uma classe conhecida.", BLUE, 22)
card(s, 8.81, 2.65, 3.85, 3.1, "Resposta OWR", "Reconhecer conhecidos, rejeitar desconhecidos e incorporar novas classes.", GOLD, 22)
banner(s, "Escopo demonstrado: um dispositivo por vez; cargas simultâneas ficam como trabalho futuro.")

s = base("02 · Arquitetura", "Um encoder; três destinos para o embedding",
         "Diagrama conceitual redesenhado a partir das seções 3.1–3.5.",
         notes="O classificador recebe o embedding, não a previsão escalar. A Figura 1 e a redação do artigo podem gerar ambiguidade; as seções 3.1 e 3.4 esclarecem o encoder compartilhado. As linhas indicam fluxo conceitual, não ordem obrigatória de execução.")
node(s, .7, 3.35, 2, 1.05, "Corrente\npré-processada", BLUE)
node(s, 3.15, 3.35, 2.05, 1.05, "Encoder\nCNN + Transformer", BLUE, 18)
node(s, 5.65, 3.35, 1.75, 1.05, "Embedding\n32 dimensões")
link(s, 2.7, 3.87, 3.15, 3.87)
link(s, 5.2, 3.87, 5.65, 3.87)
for y, title, color in [(2.45, "Cabeça Seq2Point → escalar", TEAL), (3.85, "Centroides + DP-GMM → rótulo / rejeição", GOLD), (5.25, "N decoders → sequências reconstruídas", BLUE)]:
    link(s, 7.4, 3.87, 8.05, y+.43)
    node(s, 8.05, y, 4.45, .95, title, color, 18)
text(s, .75, 5.1, 6.5, .85, "O vetor é contínuo: não é um código VQ nem um bitstream de compressão.", 21, MUTED)

s = base("02 · Arquitetura", "Seq2Point e Seq2Seq: duas tarefas de treinamento",
         "A distinção está no formato da saída, não em uma arquitetura obrigatória.",
         source=PAPER+" · §§3.3–3.4 e 4.1")
card(s, .65, 2.55, 5.85, 3.25, "Seq2Point · sequência → ponto", "Janela → embedding → um valor de corrente.\nNo protocolo final: ponto a 270° do ciclo alinhado.\nPerda auxiliar ajuda a moldar o encoder.", size=21)
card(s, 6.72, 2.55, 5.85, 3.25, "Seq2Seq · sequência → sequência", "Janela → embedding → 1.666 amostras.\nUm decoder por dispositivo conhecido.\nSó o decoder da classe da amostra recebe a perda de reconstrução.", BLUE, 21)
banner(s, "Na inferência descrita, classificação e reconstrução não dependem do escalar Seq2Point.")

s = base("02 · Arquitetura", "Da corrente medida ao sinal reconstruído",
         "A entrada efetiva do encoder tem três canais, após pré-processamento.",
         source=PAPER+" · §§3.2–3.5 e 4.1",
         notes="Passa-baixa Butterworth de quarta ordem, fase zero, corte de 500 Hz; mediana de cinco amostras; Z-score; alinhamento pela fase fundamental. Seno/cosseno preservam a fase original. O artigo não detalha se todos os decoders são executados na inferência ou apenas o selecionado. A reconstrução não recupera componentes removidos pela filtragem.")
card(s, .65, 2.55, 3.85, 3.15, "Aquisição", "1.666 amostras por janela.\n100 kHz: ≈ 16,66 ms.\nAproximadamente um ciclo de 60 Hz.", size=22)
card(s, 4.73, 2.55, 3.85, 3.15, "Pré-processamento", "Filtragem e Z-score.\nAlinhamento de fase.\nCanais: corrente alinhada, seno e cosseno da fase.", BLUE, 21)
card(s, 8.81, 2.55, 3.85, 3.15, "Saídas", "Classe ou desconhecido.\nJanela reconstruída pelo decoder correspondente.\nRotação de fase desfeita.", TEAL, 21)
banner(s, "A referência de reconstrução é o sinal pré-processado, não a corrente bruta com todo seu ruído.")

s = base("02 · Arquitetura", "Como o sistema incorpora um dispositivo novo",
         "Expansão é um processo de aprendizado; não acontece em um único forward pass.",
         source=PAPER+" · §§3.5–3.6 e 4.5")
steps = [("01 · Rejeitar", "Distância ao centroide ou baixa verossimilhança no DP-GMM."),
         ("02 · Confirmar", "Acumular janelas; nos experimentos, concordância em 8 de 10."),
         ("03 · Aprender", "Treinar apenas o novo decoder com dados do novo dispositivo."),
         ("04 · Incorporar", "Atualizar identificação; manter encoder e decoders antigos congelados.")]
for i, (heading, body) in enumerate(steps):
    card(s, .65+i*3.06, 2.6, 2.85, 3.12, heading, body, BLUE if i % 2 else TEAL, 20)
banner(s, "Congelar preserva as funções antigas; não garante representar bem todo domínio futuro.")

s = base("03 · Resultados", "A CNN pura foi melhor nesta comparação de reconstrução",
         "Mesmos decoders, MSE e condições de treino · cinco dispositivos · partição sorteada.",
         source=PAPER+" · §4.3, Tabela 2",
         notes="Valores transcritos da Tabela 2. A seleção de Deep CNN+T é justificada pelos autores pelo contexto temporal e papel no espaço de identificação/OWR. A tabela não prova ganho de identificação causado pela atenção. K significa mil e M milhão de parâmetros reportados. SDR: razão sinal-distorção; maior é melhor.")
table(s, ["Encoder", "Correlação ↑", "SDR ↑", "Parâmetros"], [
    ["Base CNN + Transformer", "0,9529", "12,52 dB", "2,68 M"],
    ["Deep CNN", "0,9999", "40,12 dB", "319 K"],
    ["Deep CNN + Transformer", "0,9998", "33,44 dB", "332 K"],
], [5.0, 2.25, 2.25, 2.2], row_h=.72, font=20)
text(s, .7, 5.65, 11.9, .42, "SDR = razão sinal–distorção. Valores maiores indicam melhor reconstrução.", 18, MUTED)
banner(s, "Transformer é uma hipótese a testar em DAS — não uma superioridade estabelecida pelo artigo.")

s = base("03 · Resultados", "Validação temporal muda a interpretação das métricas",
         "Janelas vizinhas da mesma gravação não são observações independentes.",
         source=PAPER+" · §4.2",
         notes="A primeira rodada reporta queda de SDR de 33,70 para 18,92 dB entre divisão sorteada e blocos. A deriva é associada ao período real de 1666,10 amostras, em vez de 1666. Após ancoramento: identificação de 99,9% ± 0,1 e SDR médio de 29,88 dB sob blocos. Estes números não são uma extensão direta da ablação da Tabela 2.")
card(s, .65, 2.55, 5.85, 3.15, "Risco identificado", "Divisão aleatória distribui ciclos vizinhos entre treino e validação.\nCom blocos contíguos, aparece degradação e deriva de fase.", GOLD, 23)
card(s, 6.72, 2.55, 5.85, 3.15, "Após ancoramento de fase", "99,9% ± 0,1 de identificação.\n29,88 dB de SDR médio.\nResultados reportados sob blocos contíguos.", TEAL, 23)
banner(s, "Não misturar estes resultados com a Tabela 2 como se fossem uma única ablação.")

s = base("03 · Resultados", "Expansão: resultados promissores no cenário avaliado",
         "Cinco dispositivos iniciais + três novos · 600 janelas por dispositivo.",
         source=PAPER+" · §§4.4–4.5, Tabela 3",
         notes="Tabela 3: novos dispositivos com correlação 0,99873–0,99988 e SDR 25,83–35,13 dB; 120 janelas de validação por dispositivo. Retenção é razão de correlações médias antes/depois, reportada como 100%. OWR: todos os desconhecidos avaliados foram detectados, confirmados em até dez janelas; isso não prova detecção universal de desconhecidos.")
card(s, .65, 2.6, 3.85, 3.1, "Reconstrução nova", "Correlação: 0,99873–0,99988.\nSDR: 25,83–35,13 dB.\nTrês dispositivos incorporados.", size=21)
card(s, 4.73, 2.6, 3.85, 3.1, "Retenção reportada", "100% da correlação média anterior.\nPesos antigos congelados: preservação por construção.", BLUE, 21)
card(s, 8.81, 2.6, 3.85, 3.1, "Limites", "Bancada controlada.\nUm dispositivo por vez.\nSem codec quantizado ou curva taxa–distorção.", GOLD, 21)
banner(s, "Esses resultados não preveem a fidelidade nem o bitrate alcançáveis em DAS bruto.")

s = base("04 · Aplicação em DAS", "Outra família de fibra não exige outra classe de saída",
         "No codec, o contrato continua sendo reconstruir o campo medido, não identificar um aparelho.",
         source=REPORT,
         notes="Parecer do projeto, não conclusão experimental do artigo. A distinção entre classe e domínio depende da tarefa. Novas condições podem causar forte deslocamento de distribuição, mesmo sem exigir rótulos adicionais. Número de núcleos não é número de canais espaciais DAS.")
table(s, ["Condição não vista", "Desafio provável para reconstrução"], [
    ["Multimode: outros modos", "Interferência, espectro e coerência espaço-temporal"],
    ["Multicore: outros núcleos", "Topologia e organização da entrada"],
    ["Doped: outro dopante", "Resposta, sensibilidade e distribuição de ruído"],
], [4.5, 7.2], row_h=.75, font=20)
text(s, .7, 5.8, 11.8, .5, "O mesmo codec deve atender doped, multicore, multimode e single-mode.", 20, MUTED)
banner(s, "OWR não é requisito inicial; generalização entre domínios continua sendo um problema real.")

s = base("04 · Aplicação em DAS", "O que incorporar — e o que não transportar automaticamente",
         "Adaptação seletiva: preservar a tarefa física e o custo real de transmissão.", source=REPORT)
table(s, ["Componente", "Decisão para DAS"], [
    ["Extração CNN + contexto", "Adaptar a canal × tempo; testar desenho multiescala"],
    ["Atenção temporal", "Ablação posterior, com tokens reduzidos ou por eixo"],
    ["Encoder compartilhado", "Manter; começar também com decoder compartilhado"],
    ["OWR / decoder por classe", "Não adotar como premissa do codec generalista"],
    ["Seq2Point / filtragem / fase", "Não copiar: exigem objetivo e justificativa próprios"],
], [4.35, 7.35], y=2.4, row_h=.59, font=18)
banner(s, "No codec: gargalo quantizado, payload explícito e nenhum caminho não codificado ao decoder.")

s = base("04 · Aplicação em DAS", "Por que testar uma CNN multiescala anisotrópica?",
         "Hipótese do projeto: não é um módulo pronto especificado no artigo.", source=REPORT,
         notes="Kernels ilustrativos, não especificação de arquitetura implementada. A anisotropia permite contexto temporal e espacial distintos; multiescala combina extensões diferentes. A CNN residual atual já tem hierarquia de contexto e a variante dilatada já foi testada. C precisa mudar concretamente a organização dos blocos e a redução/reconstrução por eixo. Não há promessa de ganho ou limite teórico estabelecido.")
card(s, .65, 2.55, 3.85, 3.2, "Eixos diferentes", "Tempo e posição física não são intercambiáveis.\nExemplos: kernel 1×9 temporal e 5×1 espacial.", BLUE, 22)
card(s, 4.73, 2.55, 3.85, 3.2, "Escalas diferentes", "Transientes curtos, oscilações longas e eventos com diferentes extensões espaciais.", TEAL, 22)
card(s, 8.81, 2.55, 3.85, 3.2, "Hipótese de eficiência", "Representar estrutura compartilhada entre canais sem apenas aumentar o número de índices.", GOLD, 22)
banner(s, "Respeitar contiguidade e atrasos físicos; a CNN atual já é profunda — mudar só o nome não basta.")

s = base("04 · Aplicação em DAS", "Codec candidato: adaptar a representação, manter o gargalo",
         "Esquema proposto; arquitetura detalhada e superioridade ainda não demonstradas.", source=REPORT)
labels = ["Janela DAS\ncanais × tempo", "Encoder CNN\nmultiescala", "Quantização VQ\n+ transmissão", "Decoder\ncompartilhado", "Campo DAS\nreconstruído"]
for i, label in enumerate(labels):
    x = .65+i*2.45
    node(s, x, 2.9, 2.2, 1.15, label, TEAL if i != 2 else BLUE, 18)
    if i:
        link(s, x-.25, 3.47, x, 3.47)
card(s, .65, 4.55, 5.85, 1.35, "Atenção: opcional e controlada", "", BLUE)
text(s, .85, 5.28, 5.4, .4, "Comparar C com C + atenção após reduzir tokens.", 18)
card(s, 6.72, 4.55, 5.85, 1.35, "Metadados também têm custo", "", GOLD)
text(s, 6.92, 5.28, 5.4, .4, "Normalização, shape, padding e side information.", 18)
banner(s, "Sem skips contínuos não codificados do encoder para o decoder.")

s = base("04 · Aplicação em DAS", "A/B/C: separar inicialização de arquitetura",
         "Comparação proposta; resultados históricos são referências, não controles automaticamente pareados.", source=REPORT)
table(s, ["Braço", "Arquitetura", "Inicialização e treino"], [
    ["A · transferência", "Nativa", "Offshore + fine-tuning FEBUS"],
    ["B · scratch", "Mesma nativa", "Aleatória + treino FEBUS-only"],
    ["C · nova hipótese", "Multiescala anisotrópica", "Aleatória + treino FEBUS-only"],
], [2.8, 4.05, 4.85], row_h=.75, font=20)
text(s, .7, 5.7, 11.7, .52, "O offshore original sem adaptação permanece como referência de origem.", 20, MUTED)
banner(s, "Mesmo corpus/split, objetivo e bits; oportunidade de convergência e compute explicitamente medidos.")

s = base("04 · Aplicação em DAS", "A pergunta central: melhor reconstrução por bit",
         "O histórico FEBUS já mostra ganhos fortes ao transmitir mais índices — com perda de compressão.",
         source="Relatório 37 · mesma validação do fold 1 · PROGRESS_CHECKPOINT_1",
         notes="Curva nominal pooled do relatório 37: rótulos 1×/2×/4× não são o offshore original. A taxa inclui índices e side conforme contabilização do estudo, mas não constitui codec completo com todos os custos de implantação. A geometria também muda contexto efetivo. Não inferir ganho à mesma taxa. Para aceitação: avaliar por configuração, amplitude/fase/bandas, caudas de canais, memória e latência, além de RMSE/correlação.")
table(s, ["Orçamento FEBUS", "RMSE relativo ↓", "Correlação ↑", "Compressão nominal"], [
    ["1×", "0,6805", "0,7173", "164,6×"],
    ["2×", "0,5904", "0,7962", "83,8×"],
    ["4×", "0,4900", "0,8656", "42,1×"],
], [3.1, 2.8, 2.5, 3.3], row_h=.7, font=20)
text(s, .7, 5.65, 11.9, .42, "Avaliar cada configuração e caudas por canal, além de amplitude, fase, memória e latência.", 18, MUTED)
banner(s, "Êxito: mesma fidelidade com menos bits, ou melhor fidelidade com os mesmos bits.")

s = base("Síntese", "Aproveitar a representação; não importar o problema de classes",
         "Três mensagens para orientar a próxima decisão científica.", source=REPORT+" · fonte primária: artigo, §§3–5")
for i, (heading, body) in enumerate([
    ("01 · O artigo", "Demonstra identificação e expansão em sinais individuais; não um codec DAS quantizado."),
    ("02 · A adaptação", "CNN canal–tempo multiescala como candidata; atenção opcional; OWR fora do núcleo inicial."),
    ("03 · Próxima evidência", "A/B/C com bits pareados, avaliação por configuração e confirmação independente."),
]):
    card(s, .65, 2.4+i*1.14, 11.95, 1.02, heading, "", TEAL if i != 1 else BLUE)
    text(s, 4.4, 2.58+i*1.14, 7.9, .72, body, 19)
banner(s, "Hipótese plausível ≠ ganho comprovado. Nenhum modelo final foi promovido.")

s = base("Referências e limites", "Fontes e leitura complementar",
         "Resultados do artigo e propostas do projeto são identificados separadamente nesta apresentação.",
         source="Fontes locais preservadas no projeto",
         notes="Fonte primária: Monitoramento não intrusivo de carga elétrica com detecção de dispositivos desconhecidos e aprendizado incremental, Arthur Coelho Estevão e Helder Roberto de Oliveira Rocha. Arquivo CBA_2026_paper__7079.pdf, seis páginas. SHA-256: 0dec3e0bb2281576c33ba5d6dc3114ef0dfc6d0332d693ff9b306fffe2f13003. Diagramas desta apresentação são redesenhos conceituais, não reproduções exatas de implementação.")
text(s, .7, 2.5, 11.8, 1.15, "Estevão, A. C.; Rocha, H. R. O.\nMonitoramento não intrusivo de carga elétrica com detecção de dispositivos desconhecidos e aprendizado incremental.", 23, NAVY, True)
text(s, .7, 4.02, 11.8, 1.8, "Artigo: CBA_2026_paper__7079.pdf · 6 páginas\nAvaliação crítica: progress_reports/60-CNN_TRANSFORMER_PAPER_AND_NEURAL_REDIRECTION.md\nHistórico: progress_reports/PROGRESS_CHECKPOINT_1 e relatório 37\nDirecionamento: progress_reports/FEBUS_GENERALIST_ROADMAP.md", 19, MUTED)
banner(s, "Diagramas conceituais e proposta DAS: elaboração do projeto, não resultados adicionais do artigo.")

prs.save(OUTPUT)
print(f"Created {OUTPUT} ({len(prs.slides)} slides)")
